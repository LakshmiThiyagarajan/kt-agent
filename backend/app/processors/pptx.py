from pptx import Presentation
from pathlib import Path

def extract_pptx(file_path: Path):
    prs = Presentation(file_path)
    text = ""

    for i, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += f"\n[Slide {i+1}]\n" + shape.text

    return text