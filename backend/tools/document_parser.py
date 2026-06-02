"""
tools/document_parser.py
-------------------------
Parse uploaded files into raw text.
Supports: .pdf, .docx, .txt, .md, .png, .jpg, .jpeg, .webp, .gif, .mp3, .mp4, .m4a, .wav, .webm, .ogg
"""
from __future__ import annotations

import io
from pathlib import Path

from core.logger import get_logger

logger = get_logger(__name__)

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}
AUDIO_EXTENSIONS = {".mp3", ".mp4", ".m4a", ".wav", ".webm", ".ogg", ".mpeg", ".mpga"}

IMAGE_MIME = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".gif": "image/gif",
}


class ParsedDocument:
    def __init__(self, filename: str, raw_text: str, metadata: dict):
        self.filename = filename
        self.raw_text = raw_text
        self.metadata = metadata


async def parse_document(filename: str, file_bytes: bytes) -> ParsedDocument:
    """Dispatch to the correct parser based on file extension."""
    suffix = Path(filename).suffix.lower()
    logger.info("parsing_document", filename=filename, format=suffix)

    if suffix == ".pdf":
        return _parse_pdf(filename, file_bytes)
    elif suffix == ".docx":
        return _parse_docx(filename, file_bytes)
    elif suffix in (".txt", ".md"):
        return _parse_text(filename, file_bytes)
    elif suffix in IMAGE_EXTENSIONS:
        return await _parse_image(filename, file_bytes, suffix)
    elif suffix in AUDIO_EXTENSIONS:
        return await _parse_audio(filename, file_bytes)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")


# ── Text-based parsers (sync) ─────────────────────────────────────────────────

def _parse_pdf(filename: str, data: bytes) -> ParsedDocument:
    from pypdf import PdfReader  # type: ignore
    reader = PdfReader(io.BytesIO(data))
    pages_text = [page.extract_text() or "" for page in reader.pages]
    raw_text = "\n\n".join(t.strip() for t in pages_text)
    return ParsedDocument(
        filename=filename,
        raw_text=raw_text,
        metadata={"pages": len(reader.pages), "word_count": len(raw_text.split())},
    )


def _parse_docx(filename: str, data: bytes) -> ParsedDocument:
    import docx  # type: ignore
    doc = docx.Document(io.BytesIO(data))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    raw_text = "\n\n".join(paragraphs)
    return ParsedDocument(
        filename=filename,
        raw_text=raw_text,
        metadata={"paragraphs": len(paragraphs), "word_count": len(raw_text.split())},
    )


def _parse_text(filename: str, data: bytes) -> ParsedDocument:
    raw_text = data.decode("utf-8", errors="replace").strip()
    return ParsedDocument(
        filename=filename,
        raw_text=raw_text,
        metadata={"word_count": len(raw_text.split())},
    )


# ── AI-powered parsers (async) ────────────────────────────────────────────────

async def _parse_image(filename: str, data: bytes, suffix: str) -> ParsedDocument:
    """Extract text from an image using GPT-4o Vision."""
    from services.openai_service import vision_completion

    mime = IMAGE_MIME.get(suffix, "image/jpeg")
    logger.info("parsing_image_ocr", filename=filename, mime=mime)

    text = await vision_completion(
        image_bytes=data,
        mime_type=mime,
        prompt=(
            "Extract ALL text content from this image. "
            "Include every word visible: headings, body text, tables, labels, captions, annotations. "
            "Preserve the logical reading order. Return only the extracted text — no commentary or explanation."
        ),
    )

    return ParsedDocument(
        filename=filename,
        raw_text=text.strip(),
        metadata={"format": "image_ocr", "mime_type": mime, "word_count": len(text.split())},
    )


async def _parse_audio(filename: str, data: bytes) -> ParsedDocument:
    """Transcribe audio to text using OpenAI Whisper."""
    from services.openai_service import transcribe_audio
    from openai import BadRequestError

    logger.info("parsing_audio_whisper", filename=filename)
    try:
        text = await transcribe_audio(filename, data)
    except BadRequestError as exc:
        raise ValueError(
            f"Audio file could not be transcribed. "
            f"Possible reasons: DRM-protected file, unsupported codec, or corrupted audio. "
            f"Try converting to plain MP3 or WAV first. (Detail: {exc})"
        ) from exc

    return ParsedDocument(
        filename=filename,
        raw_text=text.strip(),
        metadata={"format": "audio_transcription", "word_count": len(text.split())},
    )
